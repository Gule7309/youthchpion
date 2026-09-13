from __future__ import annotations

import re
from io import BytesIO
from typing import Any
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from app.http import RetryingHttpClient
from app.models import FreshnessStatus, SourceSnapshot, utc_now
from app.sources.base import AdapterResult

DGBAS_ANNUAL_INDEX_URL = "https://www.stat.gov.tw/News.aspx?n=4001"
TABLE_NUMBER = 47

# The public workbook is bilingual. Rows 14-20 are the seven published occupation groups;
# row names are mapped only after verifying their embedded English label.
ROW_MAP = {
    14: ("1", "主管及經理人員"),
    15: ("2", "專業人員"),
    16: ("3", "技術員及助理專業人員"),
    17: ("4", "事務支援人員"),
    18: ("5", "服務及銷售工作人員"),
    19: ("6", "農林漁牧業生產人員"),
    20: ("7-9", "技藝、機械操作及基層技術人員"),
}
EXPECTED_ENGLISH = {
    14: "Managers",
    15: "Professionals",
    16: "Technicians",
    17: "Clerical Support",
    18: "Service & Sales",
    19: "Skilled Agricultural",
    20: "Craft & Machine",
}


def _number(value: Any) -> int:
    if value in (None, "-", ""):
        return 0
    return int(round(float(value) * 1000))


def discover_latest_release(body: bytes, base_url: str) -> tuple[int, str]:
    soup = BeautifulSoup(body, "html.parser")
    releases: list[tuple[int, str]] = []
    for link in soup.find_all("a", href=True):
        title = " ".join(link.get_text(" ", strip=True).split())
        match = re.fullmatch(r"(\d{3})\s*年人力資源調查統計", title)
        if match:
            releases.append((int(match.group(1)) + 1911, urljoin(base_url, link["href"])))
    if not releases:
        raise ValueError("DGBAS annual index does not contain a human-resources release")
    return max(releases, key=lambda item: item[0])


def discover_table_url(body: bytes, base_url: str, table_number: int = TABLE_NUMBER) -> str:
    soup = BeautifulSoup(body, "html.parser")
    marker = re.compile(rf"表\s*{table_number}(?:\D|$)")
    filename = re.compile(rf"(?:^|/)table0*{table_number}\.xlsx(?:$|\?)", re.IGNORECASE)
    for link in soup.find_all("a", href=True):
        href = str(link["href"])
        if filename.search(href):
            return urljoin(base_url, href)
    for link in soup.find_all("a", href=True):
        href = str(link["href"])
        label = link.get_text(" ", strip=True).upper()
        parent_text = " ".join(link.parent.get_text(" ", strip=True).split())
        if marker.search(parent_text) and (href.lower().endswith(".xlsx") or "EXCEL" in label):
            return urljoin(base_url, href)
    raise ValueError(f"DGBAS release does not contain table {table_number} Excel")


def dgbas_workbook_period(body: bytes) -> int:
    sheet = load_workbook(BytesIO(body), data_only=True, read_only=True).active
    value = sheet.cell(5, 19).value
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("DGBAS workbook is missing the calendar year in S5") from exc


def parse_dgbas_workbook(
    body: bytes,
    expected_period: int | None = None,
) -> list[dict[str, Any]]:
    sheet = load_workbook(BytesIO(body), data_only=True).active
    data_period = dgbas_workbook_period(body)
    if expected_period is not None and data_period != expected_period:
        raise ValueError(
            f"DGBAS period mismatch: release={expected_period}, workbook={data_period}"
        )
    records: list[dict[str, Any]] = []
    for row, (code, name) in ROW_MAP.items():
        raw_label = str(sheet.cell(row, 2).value or "")
        if EXPECTED_ENGLISH[row].lower() not in raw_label.lower():
            raise ValueError(f"DGBAS schema drift at row {row}: missing English occupation label")
        # Columns: 15=20-24, 17=25-29. Values are thousands of people.
        youth_employed_20_24 = _number(sheet.cell(row, 15).value)
        youth_employed_25_29 = _number(sheet.cell(row, 17).value)
        total_employed = _number(sheet.cell(row, 3).value)
        records.append(
            {
                "code": code,
                "name": name,
                # The primary policy cohort is 20–24. Keep 25–29 separately as
                # context instead of hiding both age bands inside one total.
                "youth_employed": youth_employed_20_24,
                "youth_employed_20_24": youth_employed_20_24,
                "youth_employed_25_29": youth_employed_25_29,
                "youth_employed_20_29": youth_employed_20_24 + youth_employed_25_29,
                "total_employed": total_employed,
                "youth_employment_share": (
                    round(youth_employed_20_24 / total_employed, 4)
                    if total_employed
                    else None
                ),
                "age_columns": {"20-24": 15, "25-29": 17},
                "data_period": data_period,
            }
        )
    return records


class DgbasAdapter:
    source_id = "dgbas_employment"

    def __init__(self, http: RetryingHttpClient) -> None:
        self.http = http

    async def fetch(self) -> AdapterResult:
        index = await self.http.get(DGBAS_ANNUAL_INDEX_URL)
        data_period, release_url = discover_latest_release(index.body, index.url)
        release = await self.http.get(release_url)
        table_url = discover_table_url(release.body, release.url)
        payload = await self.http.get(table_url)
        records = parse_dgbas_workbook(payload.body, expected_period=data_period)
        snapshot = SourceSnapshot(
            source_id=self.source_id,
            status=FreshnessStatus.LIVE,
            source_url=payload.url,
            dataset_name="主計總處人力資源調查統計年報表 47",
            reference_url=release.url,
            discovery_url=index.url,
            discovery_sha256=index.sha256,
            reference_sha256=release.sha256,
            data_period=str(data_period),
            processing_steps=[
                "從官方年報清單選擇最新人力資源調查年度與表 47 Excel",
                "比對發布頁年度與工作簿 S5 西元年，不一致即停止發布",
                "驗證表 47 職業列與英文標籤，欄位漂移即停止發布",
                "擷取 20–24 歲作主要政策分析；25–29 歲保留為獨立比較欄",
                "將原始單位由千人換算為整數人數",
                "統一為七個職業大類，供 ILO 指標對齊",
            ],
            fields_used=[
                "B 欄：職業中英文名稱（schema 驗證與職類對照）",
                "O 欄：20–24 歲就業人數（主要分析）",
                "Q 欄：25–29 歲就業人數（比較脈絡）",
            ],
            why_used=(
                "建立剛進入職場的 20–24 歲青年在各職業的官方就業分布，"
                "作為風險指標的母體權重。"
            ),
            limitations="這是就業結構，不是失業率，也不能單獨證明 AI 導致失業。",
            input_count_label="個選定職業列",
            output_count_label="個 20–24 歲職業指標",
            retrieved_at=utc_now(),
            # The workbook endpoint does not expose a machine-readable publication
            # timestamp, so do not invent one. retrieved_at and the content hash
            # prove when and what this run fetched.
            source_published_at=None,
            http_status=payload.status_code,
            content_type=payload.content_type,
            content_sha256=payload.sha256,
            raw_rows=7,
            normalized_rows=len(records),
            message=(
                f"Discovered official {data_period} Table 47; release period, workbook period "
                "and schema validated after download."
            ),
        )
        return AdapterResult(
            snapshot=snapshot,
            records=records,
            audit={
                "age_filter": ["20-24", "25-29"],
                "unit_conversion": "thousand→person",
                "data_period": data_period,
                "discovery_url": index.url,
                "release_url": release.url,
                "table_url": payload.url,
            },
            raw_body=payload.body,
        )
