from __future__ import annotations

import json
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.http import HttpPayload
from app.sources.dgbas import (
    EXPECTED_ENGLISH,
    ROW_MAP,
    dgbas_workbook_period,
    discover_latest_release,
    discover_table_url,
    parse_dgbas_workbook,
)
from app.sources.dgbas_microdata import DgbasMicrodataAdapter, parse_dgbas_microdata
from app.sources.ilo import parse_ilo_csv
from app.sources.moda_public_opinion import (
    discover_latest_survey_page,
    discover_report_pdf,
    parse_public_opinion_pages,
)
from app.sources.taiwanjobs import TaiwanJobsAdapter, parse_taiwanjobs_xml
from app.sources.vacancy_history import parse_vacancy_history


def _microdata_line(*, age: int, occupation: int, weight: int) -> str:
    fields = [" "] * 91
    fields[21:24] = f"{age:03d}"
    fields[80:82] = f"{occupation:02d}"
    fields[86:91] = f"{weight:05d}"
    return "".join(fields)


def test_dgbas_microdata_uses_exact_age_boundaries_and_annual_weights() -> None:
    body = "\n".join(
        [
            _microdata_line(age=17, occupation=21, weight=1200),
            _microdata_line(age=18, occupation=21, weight=1200),
            _microdata_line(age=20, occupation=21, weight=1200),
            _microdata_line(age=24, occupation=21, weight=1200),
            _microdata_line(age=25, occupation=41, weight=1200),
            _microdata_line(age=29, occupation=41, weight=1200),
            _microdata_line(age=30, occupation=71, weight=1200),
            _microdata_line(age=35, occupation=71, weight=1200),
            _microdata_line(age=36, occupation=71, weight=1200),
            _microdata_line(age=22, occupation=1, weight=1200),
        ]
    ).encode("ascii")

    records, audit = parse_dgbas_microdata(body)

    professional = next(record for record in records if record["code"] == "2")
    clerical = next(record for record in records if record["code"] == "4")
    combined = next(record for record in records if record["code"] == "7-9")
    assert professional["total_employed"] == 400
    assert professional["youth_employed_18_35"] == 300
    assert professional["youth_employed_18_24"] == 300
    assert professional["youth_employed_20_24"] == 200
    assert clerical["youth_employed_25_29"] == 200
    assert combined["total_employed"] == 300
    assert combined["youth_employed_30_35"] == 200
    assert combined["youth_employed_18_35"] == 200
    assert audit["sample_rows_by_age_group"] == {
        "18-35": 7,
        "18-24": 3,
        "20-24": 2,
        "25-29": 2,
        "30-35": 2,
    }


def test_dgbas_microdata_fails_closed_on_invalid_layout() -> None:
    with pytest.raises(ValueError, match="expected at least 91"):
        parse_dgbas_microdata(b"too short\n")


async def test_dgbas_microdata_adapter_does_not_emit_raw_person_rows(tmp_path) -> None:
    source = tmp_path / "licensed.dat"
    source.write_text(
        "\n".join(
            _microdata_line(age=age, occupation=occupation, weight=1200)
            for age, occupation in zip(
                (18, 20, 24, 25, 29, 30, 35),
                (11, 21, 31, 41, 51, 61, 71),
                strict=True,
            )
        ),
        encoding="ascii",
    )

    result = await DgbasMicrodataAdapter(
        data_period=2024,
        local_path=source,
        minimum_rows=7,
    ).fetch()

    assert len(result.records) == 7
    assert result.raw_artifacts == []
    assert result.snapshot.data_period == "2024"
    assert result.snapshot.content_sha256


def test_dgbas_parser_extracts_20_to_29_and_converts_thousands() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(5, 19, 2025)
    for row in ROW_MAP:
        sheet.cell(row, 2, EXPECTED_ENGLISH[row])
        sheet.cell(row, 3, 100)
        sheet.cell(row, 15, 3)
        sheet.cell(row, 17, 4)
    body = BytesIO()
    workbook.save(body)

    records = parse_dgbas_workbook(body.getvalue(), expected_period=2025)

    assert len(records) == 7
    assert records[0]["youth_employed"] == 3000
    assert records[0]["youth_employed_20_24"] == 3000
    assert records[0]["youth_employed_25_29"] == 4000
    assert records[0]["youth_employed_20_29"] == 7000
    assert records[0]["total_employed"] == 100000
    assert records[0]["data_period"] == 2025
    assert dgbas_workbook_period(body.getvalue()) == 2025


def test_dgbas_parser_fails_closed_on_schema_drift() -> None:
    workbook = Workbook()
    workbook.active.cell(5, 19, 2025)
    body = BytesIO()
    workbook.save(body)

    try:
        parse_dgbas_workbook(body.getvalue())
    except ValueError as exc:
        assert "schema drift" in str(exc)
    else:
        raise AssertionError("schema drift must fail")


def test_dgbas_parser_fails_closed_on_period_mismatch() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(5, 19, 2024)
    for row in ROW_MAP:
        sheet.cell(row, 2, EXPECTED_ENGLISH[row])
    body = BytesIO()
    workbook.save(body)

    with pytest.raises(ValueError, match="period mismatch"):
        parse_dgbas_workbook(body.getvalue(), expected_period=2025)


def test_dgbas_discovery_selects_latest_release_and_table_47() -> None:
    index = """
    <a href='/release/2024'>113年人力資源調查統計</a>
    <a href='/release/2025'>114年人力資源調查統計</a>
    """.encode()
    period, release_url = discover_latest_release(index, "https://www.stat.gov.tw/list")
    release = """
    <ul><li>表46 就業者 ( <a href='/table46.xlsx'>EXCEL</a> )</li>
    <li>表47 就業者之教育程度與年齡按職業分
    ( <a href='https://ws.dgbas.gov.tw/table47.xlsx'>EXCEL</a> )</li></ul>
    """.encode()

    assert period == 2025
    assert release_url == "https://www.stat.gov.tw/release/2025"
    assert (
        discover_table_url(release, release_url)
        == "https://ws.dgbas.gov.tw/table47.xlsx"
    )


def test_dgbas_table_discovery_does_not_take_first_excel_from_large_container() -> None:
    release = b"""<ul><li>
    <p>Table 1 <a href='/files/table1.xlsx'>EXCEL</a></p>
    <p>Table 47 <a href='/files/table47.xlsx'>EXCEL</a></p>
    </li></ul>"""

    assert discover_table_url(release, "https://www.stat.gov.tw/release") == (
        "https://www.stat.gov.tw/files/table47.xlsx"
    )


def test_ilo_parser_aggregates_detailed_occupations() -> None:
    body = (
        b"Major groups,Job title,Average score,mean_exposure_level,Standard deviation\n"
        b"4 - Clerical,Clerk,0.7,Gradient 4,0.1\n"
        b"4 - Clerical,Secretary,0.5,Gradient 3,0.1\n"
        b"2 - Professionals,Engineer,0.2,Gradient 1,0.1\n"
    )

    records = parse_ilo_csv(body)

    clerical = next(item for item in records if item["code"] == "4")
    assert clerical["exposure_score"] == 0.6
    assert clerical["exposure_p90"] == 0.68
    assert clerical["occupation_count"] == 2


def test_ilo_parser_uses_distribution_not_single_highest_job_for_level() -> None:
    body = (
        b"Major groups,Job title,Average score,mean_exposure_level,Standard deviation\n"
        b"5 - Service,One,0.1,Not Exposed,0.1\n"
        b"5 - Service,Two,0.2,Not Exposed,0.1\n"
        b"5 - Service,Three,0.8,Highest exposure low task variability gradient 4,0.1\n"
    )

    service = parse_ilo_csv(body)[0]

    assert service["exposure_level"] == "Not Exposed"
    assert service["high_exposure_occupation_share"] == pytest.approx(1 / 3, abs=0.0001)


def test_moda_discovers_latest_survey_and_official_pdf() -> None:
    index = """
    <a href='/survey/2024'>113年數位近用調查報告</a>
    <a href='/survey/2025'>114年數位近用調查報告</a>
    <a href='/research/2025'>114年數位近用研究報告</a>
    """.encode()
    period, page_url = discover_latest_survey_page(index, "https://moda.gov.tw/list")
    page = """
    <a href='https://www-api.moda.gov.tw/File/Get/report'>
      114年數位近用調查報告及摘要 PDF
    </a>
    """.encode()

    assert period == 2025
    assert page_url == "https://moda.gov.tw/survey/2025"
    assert discover_report_pdf(page, page_url) == (
        "https://www-api.moda.gov.tw/File/Get/report"
    )


def test_moda_parser_cross_validates_latest_youth_opinion() -> None:
    history_rows = [[float(index)] * 5 for index in range(25)]
    history_rows[-6] = [34.0, 29.9, 31.8, 39.5, 35.0]
    history = "\n".join(
        ["表 8-4 109-114 年 20-29 歲", "構面 109 年 111 年 112 年 113 年 114 年"]
        + ["row " + " ".join(f"{value:.1f}" for value in row) for row in history_rows]
    )
    cross_rows = [[float(index)] * 10 for index in range(17)]
    cross_rows[-6] = [35.5, 54.4, 34.9, 35.9, 37.8, 35.2, 31.0, 28.6, 21.2, 15.7]
    cross = "\n".join(
        ["表 8-2 114 年 20-29 歲"]
        + ["row " + " ".join(f"{value:.1f}" for value in row) for row in cross_rows]
    )

    record, audit = parse_public_opinion_pages([history, cross], 2025)

    assert record["value"] == 35.0
    assert record["survey_year"] == 2025
    assert record["series"][-2:] == [
        {"year": 2024, "value": 39.5},
        {"year": 2025, "value": 35.0},
    ]
    assert audit["cross_table_value"] == 34.9


def test_moda_parser_fails_closed_when_tables_disagree() -> None:
    history_rows = [[1.0] * 5 for _ in range(25)]
    history_rows[-6] = [34.0, 29.9, 31.8, 39.5, 35.0]
    history = "\n".join(
        ["表 8-4 109-114 年 20-29 歲", "構面 109 年 111 年 112 年 113 年 114 年"]
        + ["row " + " ".join(f"{value:.1f}" for value in row) for row in history_rows]
    )
    cross_rows = [[1.0] * 10 for _ in range(17)]
    cross_rows[-6] = [35.5, 54.4, 30.0, 35.9, 37.8, 35.2, 31.0, 28.6, 21.2, 15.7]
    cross = "\n".join(
        ["表 8-2 114 年 20-29 歲"]
        + ["row " + " ".join(f"{value:.1f}" for value in row) for row in cross_rows]
    )

    with pytest.raises(ValueError, match="cross-table validation failed"):
        parse_public_opinion_pages([history, cross], 2025)


def test_taiwanjobs_parser_repairs_tags_filters_and_audits() -> None:
    xml = """<?xml version='1.0'?><DataList>
    <Data><OCCU_DESC（職務名稱）>AI 行政助理</OCCU_DESC（職務名稱）>
    <WK_TYPE（性質）>全職</WK_TYPE（性質）>
    <CJOB1_COUNT（大類）>01</CJOB1_COUNT（大類）><CJOB2_COUNT（小類）>010206</CJOB2_COUNT（小類）>
    <CJOB_NAME1（類別）>行政</CJOB_NAME1（類別）>
    <JOB_DETAIL（內容）>使用生成式 AI</JOB_DETAIL（內容）>
    <URL_QUERY（網址）>https://example.com/a</URL_QUERY（網址）><EXPERIENCE（經驗）>無</EXPERIENCE（經驗）>
    <JOB_PERSON（人數）>2</JOB_PERSON（人數）><STOP_DATE（截止）>額滿為止</STOP_DATE（截止）></Data>
    <Data><OCCU_DESC（職務名稱）>過期客服</OCCU_DESC（職務名稱）>
    <CJOB_NAME1（類別）>客服</CJOB_NAME1（類別）><URL_QUERY（網址）>https://example.com/b</URL_QUERY（網址）>
    <STOP_DATE（截止）>20200101</STOP_DATE（截止）></Data></DataList>"""

    records, audit = parse_taiwanjobs_xml(xml.encode(), today=date(2026, 9, 12))

    assert len(records) == 1
    assert records[0]["occupation_code"] == "4"
    assert records[0]["official_occupation_code"] == "010206"
    assert records[0]["occupation_mapping_method"] == "official_major_code_default"
    assert records[0]["quality_entry"] is True
    assert records[0]["ai_skill_category"] == "ai_application"
    assert records[0]["ai_related"] is True
    assert records[0]["headcount"] == 2
    assert audit["expired_removed"] == 1
    assert audit["schema_repairs"]


def test_taiwanjobs_distinguishes_unknown_experience_and_generic_python() -> None:
    xml = """<?xml version='1.0'?><DataList>
    <Data><OCCU_DESC（職務名稱）>Python 程式設計師</OCCU_DESC（職務名稱）>
    <WK_TYPE（性質）>全職</WK_TYPE（性質）><CJOB1_COUNT（大類）>08</CJOB1_COUNT（大類）>
    <CJOB2_COUNT（小類）>080202</CJOB2_COUNT（小類）><CJOB_NAME1（類別）>資訊</CJOB_NAME1（類別）>
    <CJOB_NAME2（小類名）>軟體設計工程師</CJOB_NAME2（小類名）>
    <JOB_DETAIL（內容）>Python 後端開發</JOB_DETAIL（內容）>
    <JOB_PERSON（人數）>3</JOB_PERSON（人數）>
    <URL_QUERY（網址）>https://example.com/python</URL_QUERY（網址）></Data>
    <Data><OCCU_DESC（職務名稱）>生成式 AI 實習生</OCCU_DESC（職務名稱）>
    <WK_TYPE（性質）>全職</WK_TYPE（性質）><CJOB1_COUNT（大類）>08</CJOB1_COUNT（大類）>
    <CJOB2_COUNT（小類）>080202</CJOB2_COUNT（小類）><CJOB_NAME1（類別）>資訊</CJOB_NAME1（類別）>
    <CJOB_NAME2（小類名）>軟體設計工程師</CJOB_NAME2（小類名）>
    <JOB_DETAIL（內容）>使用 RAG 與 LLM</JOB_DETAIL（內容）>
    <EXPERIENCE（經驗）></EXPERIENCE（經驗）>
    <JOB_PERSON（人數）>2</JOB_PERSON（人數）><URL_QUERY（網址）>https://example.com/ai</URL_QUERY（網址）></Data>
    </DataList>"""

    records, audit = parse_taiwanjobs_xml(xml.encode(), today=date(2026, 9, 12))

    python_job, ai_job = records
    assert python_job["entry_level"] is None
    assert python_job["ai_skill_category"] == "generic_programming"
    assert python_job["ai_related"] is False
    assert ai_job["entry_level"] is True
    assert ai_job["ai_skill_category"] == "ai_development"
    assert audit["entry_level_counts"] == {"true": 1, "false": 0, "unknown": 1}
    assert audit["ai_subsample_crosswalk_coverage_by_headcount"] == 1.0


async def test_taiwanjobs_adapter_queries_official_code_strata_and_keeps_raws() -> None:
    class FakeHttp:
        calls: list[str] = []

        async def get(self, url: str, params: dict | None = None) -> HttpPayload:
            jobno = str((params or {})["jobno"])
            self.calls.append(jobno)
            small_code = "010206" if jobno == "01" else "080202"
            body = f"""<?xml version='1.0'?><DataList><Data>
            <OCCU_DESC（名稱）>職缺 {jobno}</OCCU_DESC（名稱）>
            <WK_TYPE（性質）>全職</WK_TYPE（性質）>
            <CJOB1_COUNT（大類）>{jobno}</CJOB1_COUNT（大類）>
            <CJOB2_COUNT（小類）>{small_code}</CJOB2_COUNT（小類）>
            <CJOB_NAME1（大類名）>測試</CJOB_NAME1（大類名）>
            <CJOB_NAME2（小類名）>測試</CJOB_NAME2（小類名）>
            <EXPERIENCE（經驗）>無</EXPERIENCE（經驗）>
            <JOB_PERSON（人數）>1</JOB_PERSON（人數）>
            <URL_QUERY（網址）>https://example.com/{jobno}</URL_QUERY（網址）>
            </Data></DataList>""".encode()
            return HttpPayload(
                url=f"{url}?jobno={jobno}&count=10",
                status_code=200,
                content_type="application/xml",
                body=body,
            )

    http = FakeHttp()
    result = await TaiwanJobsAdapter(  # type: ignore[arg-type]
        http, count=10, jobnos=("01", "08")
    ).fetch()

    assert http.calls == ["01", "08"]
    assert len(result.records) == 2
    assert [artifact.name for artifact in result.raw_artifacts] == [
        "jobno-01",
        "jobno-08",
    ]
    assert result.audit["queried_jobnos"] == ["01", "08"]


def test_vacancy_history_builds_independent_recruitment_weakening() -> None:
    rows = []
    names = [
        "民意代表、主管及經理人員",
        "專業人員",
        "技術員及助理專業人員",
        "事務支援人員",
        "服務及銷售工作人員",
        "農、林、漁、牧業生產人員",
        "技藝有關工作人員",
        "機械設備操作及組裝人員",
        "基層技術工及勞力工及其他",
    ]
    for roc_year, value in ((111, 80), (113, 100), (114, 85)):
        rows.extend(
            {
                "統計期": f"{roc_year}年",
                "職業別": name,
                "新登記求才人數（人次）": str(value),
            }
            for name in names
        )

    records = parse_vacancy_history(json.dumps(rows, ensure_ascii=False).encode())
    clerical = next(record for record in records if record["code"] == "4")
    combined = next(record for record in records if record["code"] == "7-9")

    assert clerical["period"] == 2025
    assert clerical["yoy_change"] == -0.15
    assert clerical["recruitment_weakening"] == 0.75
    assert clerical["three_year_change"] == 0.0625
    assert combined["new_vacancies"] == 255
    assert combined["previous_new_vacancies"] == 300
